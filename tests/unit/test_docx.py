from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from bebbi_deep_research.services.docx_export import docx_filename, render_docx
from bebbi_deep_research.services.report import build_sources

from .conftest import LEDGER, REPORT

MD = """# Fasnacht im Grossen Rat

## Zusammenfassung

Der **Grosse Rat** behandelte *mehrere* Vorstösse [1]. Siehe auch [Fedlex](https://www.fedlex.admin.ch/x) und https://example.org/a.
Zweite Zeile desselben Absatzes [2]. Eine Jahreszahl [2024] ist keine Quelle.

### Ergebnisse

- Erster Punkt [1]
- Zweiter Punkt
  - Untergeordnet
1. Eins
2. Zwei

| Jahr | Anzahl | Bemerkung |
|---|:---:|---|
| 2023 | 4 [2] | **fett** |
| 2024 | 5 | a \\| b |

> Zitat aus dem Gesetz

## Quellen

[1] Vorstoss A. https://grosserrat.bs.ch/a (Stand 2023-05-01; abgerufen am 2026-10-09)
[2] Gesetz B: keine öffentliche URL (abgerufen am 2026-10-09)
"""


def render(md=MD, **kw):
    sources = build_sources(md, {"cards": []})
    data = render_docx(
        md, sources, task="Wie war die Fasnacht?", preset="grossrat", depth="quick", finished=1791559369, **kw
    )
    return Document(BytesIO(data)), sources


def texts(doc):
    return [p.text for p in doc.paragraphs]


def tables(doc):
    """The tables of the report, without the template's invisible 1pt spacer table on the first page."""
    return [t for t in doc.tables if t.rows[0].cells[0].paragraphs[0].style.name != "1pt"]


def num(p):
    """(list level, numbering id) of a list paragraph."""
    pr = p._p.pPr.find(qn("w:numPr")) if p._p.pPr is not None else None
    if pr is None:
        return None
    nid = pr.find(qn("w:numId"))
    return int(pr.find(qn("w:ilvl")).get(qn("w:val"))), None if nid is None else nid.get(qn("w:val"))


def external_links(doc):
    return {r.target_ref for r in doc.part.rels.values() if r.is_external}


def test_opens_and_has_title_and_header_block():
    doc, _ = render()
    assert [p.text for p in doc.paragraphs if p.style.name == "Title"] == ["Fasnacht im Grossen Rat"]
    head = tables(doc)[0]
    rows = {r.cells[0].text: r.cells[1].text for r in head.rows}
    assert rows["Frage"] == "Wie war die Fasnacht?" and rows["Recherchetyp"] == "Grosser Rat (grossrat)"
    assert rows["Datum"] == "09.10.2026, 17:22" and rows["Tiefe"] == "schnell (quick)"
    assert (
        doc.core_properties.title == "Fasnacht im Grossen Rat" and doc.core_properties.author == "Bebbi Deep Research"
    )
    assert texts(doc).count("Fasnacht im Grossen Rat") == 1  # the title is not repeated as heading


def test_headings_paragraphs_lists():
    doc, _ = render()
    styles = [(p.style.name, p.text) for p in doc.paragraphs]
    # heading levels shift so that the report's `##` is the template's numbered Heading 1, as in kdkp_office
    assert ("Heading 1", "Zusammenfassung") in styles and ("Heading 2", "Ergebnisse") in styles
    assert ("Heading 1", "Quellen") in styles
    assert ("ListWithBullets", "Erster Punkt [1]") in styles and ("ListWithBullets", "Zweiter Punkt") in styles
    sub = next(p for p in doc.paragraphs if p.text == "Untergeordnet")
    assert sub.style.name == "ListWithBullets" and num(sub)[0] == 1
    numbered = [p for p in doc.paragraphs if p.style.name == "ListWithNumbers"]
    assert [p.text for p in numbered] == ["Eins", "Zwei"]  # Word numbers them
    assert num(numbered[0])[1] is not None and num(numbered[0]) == num(numbered[1])
    assert sub.paragraph_format.space_after is not None  # the bullet list ends before the numbered one
    para = next(p for p in doc.paragraphs if p.text.startswith("Der Grosse Rat"))
    assert "Zweite Zeile desselben Absatzes [2]." in para.text  # soft-wrapped lines are one paragraph
    bold = [r.text for r in para.runs if r.bold]
    assert bold == ["Grosse Rat"] and [r.text for r in para.runs if r.italic] == ["mehrere"]
    assert "[2024]" in para.text
    q = next(p for p in doc.paragraphs if p.text == "Zitat aus dem Gesetz")
    assert q.style.name == "Block Text"


def test_markdown_table_becomes_word_table():
    doc, _ = render()
    assert len(tables(doc)) == 2
    t = tables(doc)[1]
    assert [[c.text for c in r.cells] for r in t.rows] == [
        ["Jahr", "Anzahl", "Bemerkung"],
        ["2023", "4 [2]", "fett"],
        ["2024", "5", "a | b"],
    ]
    assert t.rows[0].cells[0].paragraphs[0].runs[0].bold
    assert t.rows[1].cells[2].paragraphs[0].runs[0].bold
    assert t.style.name == "Table Grid"
    assert [r.cells[1].paragraphs[0].alignment for r in t.rows] == [WD_ALIGN_PARAGRAPH.CENTER] * 3  # `:---:`
    assert t.rows[1].cells[0].paragraphs[0].alignment is None


def test_citations_link_to_sources_and_source_list_has_hyperlinks():
    doc, sources = render()
    body = doc.element.body
    anchors = [h.get(qn("w:anchor")) for h in body.iter(qn("w:hyperlink")) if h.get(qn("w:anchor"))]
    assert anchors.count("quelle_1") >= 2 and "quelle_2" in anchors and "quelle_2024" not in anchors
    bookmarks = {b.get(qn("w:name")) for b in body.iter(qn("w:bookmarkStart"))}
    assert bookmarks == {"quelle_1", "quelle_2"}
    assert external_links(doc) == {"https://grosserrat.bs.ch/a"}  # only registered source URLs (see below)
    listing = [t for t in texts(doc) if t.startswith("[1]\t")]
    assert len(listing) == 1 and "Vorstoss A" in listing[0] and "https://grosserrat.bs.ch/a" in listing[0]
    assert "Stand 2023-05-01; abgerufen am 2026-10-09; Art der Quelle: Web" in listing[0]
    nourl = next(t for t in texts(doc) if t.startswith("[2]\t"))
    assert "keine öffentliche URL" in nourl and "Quellen" in texts(doc)
    assert sum(1 for t in texts(doc) if "abgerufen am" in t) == 2  # the Markdown list is not repeated


def test_footer_and_swiss_spelling():
    doc, _ = render()
    for section in doc.sections:  # first page and following pages
        footer = section.footer.paragraphs[0]
        assert footer.text.startswith("Erstellt mit Bebbi Deep Research · Seite ")
        assert "PAGE" in footer._p.xml and "NUMPAGES" in footer._p.xml
    assert "ß" not in "\n".join(texts(doc))


def test_report_without_source_list_is_rendered_as_is():
    md = "# T\n\nText mit [1] und\n\n## Quellen\n\n[1] kaputt"
    doc = Document(BytesIO(render_docx(md, [], task="x", preset="p", depth="deep")))
    assert "[1] kaputt" in texts(doc) and tables(doc)[0].rows[3].cells[1].text == "gründlich (deep)"


def test_english_report_and_fake_agent_report():
    md = "# T\n\nText [1].\n\n## Sources\n\n[1] Title. https://a.org (dated 2021-01-01; accessed 2026-10-09)\n"
    doc, _ = render(md)
    assert tables(doc)[0].rows[0].cells[0].text == "Question" and "dated 2021-01-01; accessed 2026-10-09" in "\n".join(
        texts(doc)
    )
    real = Document(
        BytesIO(render_docx(REPORT, build_sources(REPORT, LEDGER), task="Frage", preset="grossrat", depth="quick"))
    )
    assert len(tables(real)) == 2 and tables(real)[1].rows[1].cells[1].text == "4 [2]"
    assert (
        "Page " in doc.sections[0].footer.paragraphs[0].text
        and doc.styles["Normal"].element.rPr.find(qn("w:lang")).get(qn("w:val")) == "en-GB"
    )


def test_filename_is_ascii():
    assert (
        docx_filename("Fasnacht: Übersicht über Grossräte (2024)", 1791559369)
        == "Bebbi-Deep-Research_Fasnacht-Uebersicht-ueber-Grossraete-2024_2026-10-09.docx"
    )
    assert docx_filename("???").startswith("Bebbi-Deep-Research_Bericht_")


def test_control_characters_do_not_break_the_export():
    report = REPORT.replace("Der Grosse Rat", "Der Gros\x1bse \x00Rat").replace("Stadtklima.", "Stadt\x0bklima.")
    sources = build_sources(report, LEDGER)
    sources[0]["excerpt"] = "Aus\x08zug"
    data = render_docx(report, sources, task="Fra\x1bge", preset="grossrat", depth="quick", finished=1760000000.0)
    text = "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)
    assert "Der Grosse Rat" in text and "\x1b" not in text


def test_only_registered_source_urls_become_links():
    md = MD.replace(
        "Siehe auch [Fedlex]",
        "Siehe [§ 12 Personalgesetz](https://evil.example/login), [Vorstoss A](https://grosserrat.bs.ch/a), "
        "<https://evil.example/auto> und [Fedlex]",
    )
    doc, _ = render(md)
    assert external_links(doc) == {"https://grosserrat.bs.ch/a"}
    para = next(p for p in doc.paragraphs if p.text.startswith("Der Grosse Rat"))
    assert "Siehe § 12 Personalgesetz, Vorstoss A, <https://evil.example/auto> und Fedlex" in para.text
    assert "https://example.org/a" in para.text  # a bare URL that is not a source stays as plain text
    link_texts = ["".join(t.text for t in h.iter(qn("w:t"))) for h in para._p.iter(qn("w:hyperlink"))]
    assert "Vorstoss A" in link_texts and "§ 12 Personalgesetz" not in link_texts


def test_report_without_sources_has_no_external_links():
    doc = Document(BytesIO(render_docx("# T\n\n[Hier](https://evil.example/x) und https://evil.example/y", [])))
    assert external_links(doc) == set() and "Hier und https://evil.example/y" in texts(doc)


def test_kanton_basel_stadt_layout_as_in_kdkp_office():
    """Template, style names and colours are those of kdkp_office.word (Bebbi Bot's document tools)."""
    doc, _ = render()
    assert doc.styles["Normal"].font.name == "Arial"
    assert len(doc.sections[0].header._element.xpath(".//w:drawing")) == 1  # the Baselstab logo on the first page
    assert round(doc.sections[0].left_margin.cm, 1) == 2.2 and round(doc.sections[1].top_margin.cm, 1) == 2.5
    styles = {p.style.name for p in doc.paragraphs}
    assert not styles & {"List Bullet", "List Bullet 2", "Heading 3"}
    body = doc.element.body
    assert not list(body.iter(qn("w:color"))) and not [
        f for f in body.iter(qn("w:rFonts")) if f.get(qn("w:ascii")) == "Calibri"
    ]
    link_styles = {
        r.find(f"{qn('w:rPr')}/{qn('w:rStyle')}").get(qn("w:val"))
        for r in body.iter(qn("w:r"))
        if r.getparent().tag == qn("w:hyperlink")
    }
    assert link_styles == {doc.styles["Hyperlink"].style_id}
    head = tables(doc)[1].rows[0].cells[0]._tc
    assert head.find(f"{qn('w:tcPr')}/{qn('w:shd')}").get(qn("w:fill")) == "E7E6E6"
    assert doc.core_properties.author == "Bebbi Deep Research"


def test_each_numbered_list_counts_from_its_first_number():
    md = "# T\n\nA\n\n1. Eins\n2. Zwei\n\nB\n\n3. Drei\n\nC\n\n1. Neu\n"
    doc = Document(BytesIO(render_docx(md, [])))
    items = [p for p in doc.paragraphs if p.style.name == "ListWithNumbers"]
    ids = [num(p)[1] for p in items]
    assert [p.text for p in items] == ["Eins", "Zwei", "Drei", "Neu"] and ids[0] == ids[1] and len(set(ids)) == 3
    numbering = doc.part.numbering_part.element
    starts = {
        n.get(qn("w:numId")): n.find(f"{qn('w:lvlOverride')}/{qn('w:startOverride')}").get(qn("w:val"))
        for n in numbering.findall(qn("w:num"))
        if n.get(qn("w:numId")) in ids
    }
    assert [starts[i] for i in ids] == ["1", "1", "3", "1"]


def test_code_block_uses_the_template_style():
    doc = Document(BytesIO(render_docx("# T\n\n```\na = 1\nb = 2\n```\n\nText", [])))
    code = [p for p in doc.paragraphs if p.style.name == "Plain Text"]
    assert [p.text for p in code] == ["a = 1", "b = 2"] and code[0].runs[0].font.name == "Consolas"
