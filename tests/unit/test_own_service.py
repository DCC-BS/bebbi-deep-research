"""Own sources in the job service: limits, storage with the job, presets (prefer / only / off), the retriever for
Open WebUI knowledge bases (only approved ids), schema migration, labels."""

import json
import re
from pathlib import Path

import httpx
import pytest
import yaml

from bebbi_deep_research.engine.own import OwnCollection
from bebbi_deep_research.models.own_sources import OwnSourcesRequest
from bebbi_deep_research.services.db import DB
from bebbi_deep_research.services.own_sources import OwuiRetriever, prepare
from bebbi_deep_research.services.report import owui_source_events, source_label

from .conftest import FakeAgent, make_settings, running, utc_window, wait_status

ROOT = Path(__file__).resolve().parents[2]


def own_payload(**kw):
    base = {
        "files": [
            {
                "id": "f1",
                "name": "merkblatt.pdf",
                "text": "Frist 28 Tage. " * 50,
                "date": "2026-10-09",
                "url": "http://owui/api/v1/files/f1/content",
            }
        ],
        "collections": [{"id": "kb1", "name": "Beschaffung"}],
        "base_url": "http://owui",
        "notes": [],
    }
    base.update(kw)
    return base


# ---------------------------------------------------------------------------------------------------- limits
def test_prepare_applies_limits_and_reports_them(tmp_path):
    s = make_settings(tmp_path, own_max_files=2, own_max_file_chars=100, own_max_collections=1)
    own = OwnSourcesRequest(
        files=[
            {"id": "a", "name": "lang.pdf", "text": "x" * 500},
            {"id": "b", "name": "b.txt", "text": "kurz"},
            {"id": "c", "name": "c.txt", "text": "drittes"},
            {"id": "a", "name": "lang.pdf", "text": "x"},
        ],
        collections=[{"id": "k1", "name": "Eins"}, {"id": "k2", "name": "Zwei"}],
    )
    stored, info = prepare(own, s)
    assert [f["id"] for f in stored["files"]] == ["a", "b"] and len(
        stored["files"][0]["text"]
    ) == 100  # cut, duplicate dropped
    assert [c["id"] for c in stored["collections"]] == ["k1"]
    notes = " ".join(info["notes"])
    assert "lang.pdf" in notes and "gekürzt" in notes and "höchstens 2 Dateien" in notes and "Zwei" in notes
    assert info["files"][0] == {"name": "lang.pdf", "chars": 500, "cut": True}
    assert "text" not in json.dumps(info)  # the user-facing view never carries file contents


def test_prepare_preview_keeps_short_text_and_empty_files_are_dropped_for_the_job(tmp_path):
    s = make_settings(tmp_path)
    own = OwnSourcesRequest(
        files=[
            {"id": "a", "name": "a.pdf", "text": "y" * 5000, "chars": 90000},
            {"id": "e", "name": "scan.pdf", "text": " "},
        ]
    )
    stored, info = prepare(own, s, preview=True)
    assert len(stored["files"][0]["text"]) == 1500 and stored["files"][0]["chars"] == 90000
    stored, info = prepare(own, s, preview=False)
    assert [f["id"] for f in stored["files"]] == ["a"] and "scan.pdf" in " ".join(info["notes"])
    assert prepare(None, s) == (None, None) and prepare(OwnSourcesRequest(), s) == (None, None)
    only_notes = prepare(OwnSourcesRequest(notes=["Datei „x“ ist für Sie nicht lesbar"]), s)
    assert only_notes[0] is None and only_notes[1]["notes"] == ["Datei „x“ ist für Sie nicht lesbar"]


def test_base_url_must_be_http(tmp_path):
    stored, _ = prepare(
        OwnSourcesRequest(collections=[{"id": "k", "name": "K"}], base_url="javascript:alert(1)"),
        make_settings(tmp_path),
    )
    assert stored["base_url"] == ""


# ----------------------------------------------------------------------------------- storage and the job
async def test_own_sources_are_stored_with_the_job_and_reach_the_engine(tmp_path):
    async with running(tmp_path, night_window=utc_window(True)) as (c, app):
        plan = (
            await c.post(
                "/plans", json={"task": "Frage", "preset": "grossrat", "depth": "quick", "own_sources": own_payload()}
            )
        ).json()
        assert plan["own_sources"]["files"][0]["name"] == "merkblatt.pdf" and plan["own_sources"]["mode"] == "prefer"
        kind, own = FakeAgent.own_seen[-1]
        assert (
            kind == "plan"
            and [f.name for f in own.files] == ["merkblatt.pdf"]
            and own.collections[0].name == "Beschaffung"
        )
        r = await c.post(
            "/jobs", json={"draft_id": plan["draft_id"], "user_id": "u1", "when": "now", "own_sources": own_payload()}
        )
        assert r.status_code == 200 and r.json()["own_sources"]["collections"] == [{"name": "Beschaffung"}]
        job = await wait_status(c, r.json()["job_id"], {"done"})
        assert job["own_sources"]["files"][0]["chars"] == len(("Frist 28 Tage. " * 50).strip())
        assert "Frist 28 Tage" not in json.dumps(job)  # the API never returns texts
        # the texts were needed by the run only: the finished job keeps names, sizes and the knowledge base ids
        stored = json.loads((await app.state.container.db().run("select own_sources from jobs"))[0]["own_sources"])
        assert (
            "text" not in stored["files"][0] and stored["files"][0]["chars"] == job["own_sources"]["files"][0]["chars"]
        )
        assert stored["collections"] == [{"id": "kb1", "name": "Beschaffung"}]
        kind, own = FakeAgent.own_seen[-1]
        assert kind == "run" and own.files[0].text.startswith("Frist") and own.files[0].url.endswith("/f1/content")
        assert own.retrieve is not None and own.collections[0].id == "kb1"


async def test_job_without_own_sources_is_unchanged(tmp_path):
    async with running(tmp_path, night_window=utc_window(True)) as (c, app):
        j = await wait_status(c, (await _submit(c))["job_id"], {"done"})
        assert j["own_sources"] is None
        assert FakeAgent.own_seen[-1] == ("run", None)


async def _submit(c, preset="grossrat", **extra):
    d = (await c.post("/plans", json={"task": "Frage", "preset": preset, "depth": "quick"})).json()
    return (await c.post("/jobs", json={"draft_id": d["draft_id"], "user_id": "u1", "when": "now", **extra})).json()


async def test_clarify_knows_the_attached_sources(tmp_path):
    async with running(tmp_path) as (c, app):
        r = await c.post("/clarify", json={"task": "Frage", "preset": "grossrat", "own_sources": own_payload()})
        assert r.status_code == 200
        kind, own = FakeAgent.own_seen[-1]
        assert kind == "clarify" and own.files[0].name == "merkblatt.pdf"


async def test_plan_edit_keeps_the_own_sources_info(tmp_path):
    async with running(tmp_path) as (c, app):
        d = (
            await c.post(
                "/plans", json={"task": "F", "preset": "grossrat", "depth": "quick", "own_sources": own_payload()}
            )
        ).json()
        e = (await c.post(f"/plans/{d['draft_id']}/edit", json={"instruction": "kürzer"})).json()
        assert e["own_sources"]["files"][0]["name"] == "merkblatt.pdf" and e["own_sources"]["mode"] == "prefer"


async def test_preset_only_needs_sources_and_off_ignores_them(tmp_path):
    cfg = yaml.safe_load((ROOT / "config/presets.yaml").read_text())
    cfg["presets"]["grossrat"]["own_sources"] = "off"
    path = tmp_path / "presets.yaml"
    path.write_text(yaml.safe_dump(cfg, allow_unicode=True))
    async with running(tmp_path, presets_path=path) as (c, app):
        r = await c.post(
            "/plans", json={"task": "F", "preset": "grossrat", "depth": "quick", "own_sources": own_payload()}
        )
        assert r.json()["own_sources"]["files"] == [] and "keine eigenen Quellen" in r.json()["own_sources"]["notes"][0]
        assert FakeAgent.own_seen[-1] == ("plan", None)  # the engine never sees them
        r = await c.post("/jobs", json={"draft_id": r.json()["draft_id"], "user_id": "u", "own_sources": own_payload()})
        assert (await app.state.container.db().run("select own_sources from jobs"))[0]["own_sources"] is None
        # `only` (the shipped preset "dokumente") refuses a question without any own source, with a message for the user
        r = await c.post("/plans", json={"task": "F", "preset": "dokumente", "depth": "quick"})
        assert r.status_code == 422 and "nur mit eigenen Quellen" in r.json()["debugMessage"]
        r = await c.post("/clarify", json={"task": "F", "preset": "dokumente"})
        assert r.status_code == 422
        r = await c.post(
            "/plans", json={"task": "F", "preset": "dokumente", "depth": "quick", "own_sources": own_payload()}
        )
        assert r.status_code == 200 and r.json()["own_sources"]["mode"] == "only"


async def test_only_preset_with_only_unreadable_sources_names_the_reason(tmp_path):
    async with running(tmp_path) as (c, app):
        r = await c.post(
            "/plans",
            json={
                "task": "F",
                "preset": "dokumente",
                "depth": "quick",
                "own_sources": {"notes": ["Datei „x.pdf“ ist für Sie nicht lesbar und wird nicht verwendet."]},
            },
        )
        assert r.status_code == 422 and "x.pdf" in r.json()["debugMessage"]


# ------------------------------------------------------------------------------------- migration, config
async def test_old_database_gets_the_new_columns(tmp_path):
    db = DB(f"sqlite+aiosqlite:///{tmp_path}/old.db")
    await db.run(
        "create table drafts (id text primary key, task text, preset text, depth text, plan text, created double precision)"
    )
    await db.run(
        "create table jobs (id text primary key, draft text, task text, preset text, depth text, plan text, user_id text,"
        " chat_id text, message_id text, run_when text, status text, created double precision, started double precision,"
        " finished double precision, events text, report text, stats text, evidence text, error text,"
        " heartbeat double precision, delivery text)"
    )
    await db.run("insert into jobs (id, status) values ('old', 'done')")
    await db.init()
    await db.init()  # idempotent
    assert (await db.run("select id, own_sources from jobs")) == [{"id": "old", "own_sources": None}]
    await db.run("select own_sources from drafts")
    await db.close()


def test_settings_for_own_source_limits():
    from bebbi_deep_research.utils.app_config import AppConfig

    required = {
        "RA_TOKEN": "t",
        "HMAC_SECRET": "h",
        "RA_LLM_BASE_URL": "u",
        "RA_SEARCH_URL": "u",
        "RA_FIRECRAWL_URL": "u",
    }
    limits = {"RA_OWN_MAX_FILES": "3", "RA_OWN_MAX_FILE_CHARS": "5000", "RA_OWN_MAX_COLLECTIONS": "2"}
    s = AppConfig.from_env({**required, **limits})
    assert (s.own_max_files, s.own_max_file_chars, s.own_max_collections) == (3, 5000, 2)
    d = AppConfig.from_env(required)
    assert (d.own_max_files, d.own_max_file_chars, d.own_max_collections) == (10, 2_000_000, 5)


# --------------------------------------------------------------------------------- retrieval (knowledge)
def _retriever(tmp_path, handler, allowed=("kb1",)):
    s = make_settings(tmp_path, owui_url="http://owui:8080", owui_api_key="admin-key")
    return OwuiRetriever(
        s,
        list(allowed),
        "https://bebbi.example",
        http_factory=lambda: httpx.AsyncClient(transport=httpx.MockTransport(handler)),
    )


async def test_retriever_queries_only_approved_collections_and_maps_hits(tmp_path):
    calls = []

    def handler(request: httpx.Request):
        calls.append(
            (request.method, request.url.path, json.loads(request.content or b"null"), request.headers["authorization"])
        )
        if request.url.path == "/api/v1/retrieval/query/collection":
            return httpx.Response(
                200,
                json={
                    "documents": [["Direktvergaben bis CHF 15'000.", "Zweiter Abschnitt"]],
                    "metadatas": [[{"file_id": "f9", "name": "Richtlinie.txt"}, {"name": "ohne-id.txt"}]],
                    "distances": [[0.68, 0.41]],
                },
            )
        return httpx.Response(200, json={"created_at": 1791571774})  # GET /api/v1/files/f9

    r = _retriever(tmp_path, handler)
    hits = await r("Direktvergabe", OwnCollection("kb1", "Beschaffung"))
    assert hits[0] == {
        "text": "Direktvergaben bis CHF 15'000.",
        "name": "Richtlinie.txt",
        "file_id": "f9",
        "url": "https://bebbi.example/api/v1/files/f9/content",
        "date": "2026-10-09",
        "score": 0.68,
    }
    assert hits[1]["url"] == "" and hits[1]["date"] == ""  # no file id: no link, nothing invented
    post = next(c for c in calls if c[1].endswith("/query/collection"))
    assert (
        post[2]["collection_names"] == ["kb1"] and post[2]["query"] == "Direktvergabe" and post[3] == "Bearer admin-key"
    )
    with pytest.raises(PermissionError):  # an id that was not approved for the job
        await r("x", OwnCollection("kb-of-someone-else", "Fremd"))
    assert all("someone-else" not in json.dumps(c[2]) for c in calls)
    assert len([c for c in calls if c[1].endswith("/query/collection")]) == 1


async def test_retriever_errors_are_raised_for_the_tool_to_report(tmp_path):
    r = _retriever(tmp_path, lambda request: httpx.Response(403, json={"detail": "no"}))
    with pytest.raises(httpx.HTTPStatusError):
        await r("x", OwnCollection("kb1", "K"))
    s = make_settings(tmp_path)  # no OWUI key configured
    with pytest.raises(RuntimeError):
        await OwuiRetriever(s, ["kb1"], "")("x", OwnCollection("kb1", "K"))


# ------------------------------------------------------------------------------------------------ labels
def test_labels_for_own_sources():
    assert source_label("upload") == "Hochgeladene Datei" and source_label("upload", de=False) == "Uploaded file"
    assert source_label("knowledge:Beschaffung und Personal") == "Wissensbasis Beschaffung und Personal"
    assert source_label("knowledge") == "Wissensbasis"
    for kind in ("upload", "knowledge:X", "knowledge"):
        assert "mcp" not in source_label(kind).lower()


def test_source_events_and_word_file_use_the_labels():
    sources = [
        {
            "n": 1,
            "title": "merkblatt.pdf",
            "url": "http://owui/api/v1/files/f1/content",
            "stand": "2026-10-09",
            "fetched": "2026-10-09",
            "kind": "upload",
            "excerpt": "Frist 28 Tage",
            "excerpt_kind": "claim",
        },
        {
            "n": 2,
            "title": "Weisung.txt",
            "url": "",
            "stand": "",
            "fetched": "2026-10-09",
            "kind": "knowledge:Personal",
            "excerpt": "",
            "excerpt_kind": "claim",
        },
    ]
    ev = owui_source_events(sources)
    assert (
        "Quelle: Hochgeladene Datei" in ev[0]["data"]["document"][0]
        and "Quelle: Wissensbasis Personal" in ev[1]["data"]["document"][0]
    )
    assert ev[1]["data"]["metadata"][0]["source"] == "quelle-2"  # no URL: a stable id, never an invented link
    import io

    from docx import Document

    from bebbi_deep_research.services.docx_export import render_docx

    report = "# T\n\n## Teil\n\nFrist [1]. Homeoffice [2].\n\n## Quellen\n\n[1] merkblatt.pdf. http://owui/api/v1/files/f1/content (Stand 2026-10-09; abgerufen am 2026-10-09)\n[2] Weisung.txt. keine öffentliche URL (abgerufen am 2026-10-09)\n"
    data = render_docx(report, sources, task="Frage", preset="dokumente", depth="quick", finished=0, tz="UTC")
    doc = Document(io.BytesIO(data))
    text = "\n".join(
        [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells]
    )
    assert "Hochgeladene Datei" in text and "Wissensbasis Personal" in text and "mcp" not in text.lower()
    assert re.search(r"Eigene Dokumente", text)


async def test_clarify_suggests_a_preset_and_sees_sources_without_a_preset_rule(tmp_path):
    async with running(tmp_path) as (c, app):
        # Before a preset is chosen the model sees the attached sources (no preset's own-sources rule applies).
        r = await c.post(
            "/clarify",
            json={"task": "Frage", "suggest_preset": True, "preset": "web_survey", "own_sources": own_payload()},
        )
        assert r.status_code == 200 and r.json()["preset"] == "grossrat"
        kind, own = FakeAgent.own_seen[-1]
        assert kind == "clarify" and own.files[0].name == "merkblatt.pdf"
        assert "preset" not in (await c.post("/clarify", json={"task": "Frage"})).json()
