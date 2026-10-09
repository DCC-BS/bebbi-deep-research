from io import BytesIO

from docx import Document
from docx.oxml.ns import qn

from service.docx_export import docx_filename, render_docx
from service.report import build_sources

from .conftest import LEDGER, REPORT

MD = """# Fasnacht im Grossen Rat

## Zusammenfassung

Der **Grosse Rat** behandelte *mehrere* Vorstösse [1]. Siehe auch [Fedlex](https://www.fedlex.admin.ch/x) und https://example.org/a.
Zweite Zeile desselben Absatzes [2]. Eine Jahreszahl [2024] ist keine Quelle.

### Ergebnisse

- Erster Punkt [1]
- Zweiter Punkt
  - Untergeordnet
1. Eins
2. Zwei

| Jahr | Anzahl | Bemerkung |
|---|:---:|---|
| 2023 | 4 [2] | **fett** |
| 2024 | 5 | a \\| b |

> Zitat aus dem Gesetz

## Quellen

[1] Vorstoss A. https://grosserrat.bs.ch/a (Stand 2023-05-01; abgerufen am 2026-10-09)
[2] Gesetz B: keine öffentliche URL (abgerufen am 2026-10-09)
"""


def render(md=MD, **kw):
    sources = build_sources(md, {"cards": []})
    data = render_docx(md, sources, task="Wie war die Fasnacht?", preset="grossrat", depth="quick", finished=1791559369, **kw)
    return Document(BytesIO(data)), sources


def texts(doc):
    return [p.text for p in doc.paragraphs]


def external_links(doc):
    return {r.target_ref for r in doc.part.rels.values() if r.is_external}


def test_opens_and_has_title_and_header_block():
    doc, _ = render()
    assert doc.paragraphs[0].style.name == "Title" and doc.paragraphs[0].text == "Fasnacht im Grossen Rat"
    head = doc.tables[0]
    rows = {r.cells[0].text: r.cells[1].text for r in head.rows}
    assert rows["Frage"] == "Wie war die Fasnacht?" and rows["Recherchetyp"] == "Grosser Rat (grossrat)"
    assert rows["Datum"] == "09.10.2026, 17:22" and rows["Tiefe"] == "schnell (quick)"
    assert doc.core_properties.title == "Fasnacht im Grossen Rat" and doc.core_properties.author == "Bebbi Deep Research"
    assert texts(doc).count("Fasnacht im Grossen Rat") == 1                    # the title is not repeated as heading


def test_headings_paragraphs_lists():
    doc, _ = render()
    styles = [(p.style.name, p.text) for p in doc.paragraphs]
    assert ("Heading 2", "Zusammenfassung") in styles and ("Heading 3", "Ergebnisse") in styles
    assert ("List Bullet", "Erster Punkt [1]") in styles and ("List Bullet", "Zweiter Punkt") in styles
    assert ("List Bullet 2", "Untergeordnet") in styles
    assert [t for _, t in styles if t.startswith(("1.", "2."))] == ["1.\tEins", "2.\tZwei"]
    para = next(p for p in doc.paragraphs if p.text.startswith("Der Grosse Rat"))
    assert "Zweite Zeile desselben Absatzes [2]." in para.text                      # soft-wrapped lines are one paragraph
    bold = [r.text for r in para.runs if r.bold]
    assert bold == ["Grosse Rat"] and [r.text for r in para.runs if r.italic] == ["mehrere"]
    assert "[2024]" in para.text
    q = next(p for p in doc.paragraphs if p.text == "Zitat aus dem Gesetz")
    assert q.runs[0].italic


def test_markdown_table_becomes_word_table():
    doc, _ = render()
    assert len(doc.tables) == 2
    t = doc.tables[1]
    assert [[c.text for c in r.cells] for r in t.rows] == [["Jahr", "Anzahl", "Bemerkung"], ["2023", "4 [2]", "fett"], ["2024", "5", "a | b"]]
    assert t.rows[0].cells[0].paragraphs[0].runs[0].bold
    assert t.rows[1].cells[2].paragraphs[0].runs[0].bold
    assert t.style.name == "Table Grid"


def test_citations_link_to_sources_and_source_list_has_hyperlinks():
    doc, sources = render()
    body = doc.element.body
    anchors = [h.get(qn("w:anchor")) for h in body.iter(qn("w:hyperlink")) if h.get(qn("w:anchor"))]
    assert anchors.count("quelle_1") >= 2 and "quelle_2" in anchors and "quelle_2024" not in anchors
    bookmarks = {b.get(qn("w:name")) for b in body.iter(qn("w:bookmarkStart"))}
    assert bookmarks == {"quelle_1", "quelle_2"}
    assert {"https://grosserrat.bs.ch/a", "https://www.fedlex.admin.ch/x", "https://example.org/a"} <= external_links(doc)
    listing = [t for t in texts(doc) if t.startswith("[1]\t")]
    assert len(listing) == 1 and "Vorstoss A" in listing[0] and "https://grosserrat.bs.ch/a" in listing[0]
    assert "Stand 2023-05-01; abgerufen am 2026-10-09; Art der Quelle: Web" in listing[0]
    nourl = next(t for t in texts(doc) if t.startswith("[2]\t"))
    assert "keine öffentliche URL" in nourl and "Quellen" in texts(doc)
    assert sum(1 for t in texts(doc) if "abgerufen am" in t) == 2                   # the Markdown list is not repeated


def test_footer_and_swiss_spelling():
    doc, _ = render()
    footer = doc.sections[0].footer.paragraphs[0]
    assert "Erstellt mit Bebbi Deep Research" in footer.text
    assert "PAGE" in footer._p.xml
    assert "ß" not in "\n".join(texts(doc))


def test_report_without_source_list_is_rendered_as_is():
    md = "# T\n\nText mit [1] und\n\n## Quellen\n\n[1] kaputt"
    doc = Document(BytesIO(render_docx(md, [], task="x", preset="p", depth="deep")))
    assert "[1] kaputt" in texts(doc) and doc.tables[0].rows[3].cells[1].text == "gründlich (deep)"


def test_english_report_and_fake_agent_report():
    md = "# T\n\nText [1].\n\n## Sources\n\n[1] Title. https://a.org (dated 2021-01-01; accessed 2026-10-09)\n"
    doc, _ = render(md)
    assert doc.tables[0].rows[0].cells[0].text == "Question" and "dated 2021-01-01; accessed 2026-10-09" in "\n".join(texts(doc))
    real = Document(BytesIO(render_docx(REPORT, build_sources(REPORT, LEDGER), task="Frage", preset="grossrat", depth="quick")))
    assert len(real.tables) == 2 and real.tables[1].rows[1].cells[1].text == "4 [2]"


def test_filename_is_ascii():
    assert docx_filename("Fasnacht: Übersicht über Grossräte (2024)", 1791559369) == "Bebbi-Deep-Research_Fasnacht-Uebersicht-ueber-Grossraete-2024_2026-10-09.docx"
    assert docx_filename("???") .startswith("Bebbi-Deep-Research_Bericht_")
